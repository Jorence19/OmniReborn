"""Private, crash-safe Telegram delivery service for OmniReborn Phase 1."""
from __future__ import annotations
import argparse
import html
import json
import logging
import os
import sqlite3
import subprocess
import sys
import tempfile
import time
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote

import requests

from database import ForensicDatabase
from source_intake import forwarded_text, is_forwarded_message, parse_forwarded_tokens
from sources import (
    SOURCE_DEFINITIONS,
    is_source_enabled,
    load_source_switches,
    normalize_source_name,
    set_source_switch,
)
from streamer import QueueStore

ROOT = Path(__file__).resolve().parent
LOG = logging.getLogger("omnireborn.telegram")
CHAINS = {4663: "RBH", 5042: "ARC"}
EXPLORERS = {
    4663: "https://robinhoodchain.blockscout.com/address/{ca}",
    5042: "https://explorer.arc.io/address/{ca}",
}


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean(value, default=""):
    result = "" if value is None else str(value).strip()
    return default if result.lower() in {"", "none", "nan"} else result


def atomic_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
        json.dump(value, handle, indent=2, sort_keys=True)
        handle.flush()
        os.fsync(handle.fileno())
        temporary = Path(handle.name)
    os.replace(temporary, path)


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@dataclass(frozen=True)
class Settings:
    token: str
    chats: tuple[int, ...]
    db: Path
    snapshot: Path
    report: Path
    csv: Path
    dashboard: Path
    runtime: Path
    health: Path
    state: Path
    lock: Path
    threshold: float = 65.0
    top: int = 5
    poll: int = 25
    push_alerts: bool = False
    alert_existing: bool = False
    max_age: int = 180

    @classmethod
    def from_env(cls, require_telegram=True):
        runtime = Path(os.getenv("RUNTIME_DIR", ROOT / "runtime")).resolve()
        token = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
        try:
            chats = tuple(
                int(item.strip())
                for item in os.getenv("TELEGRAM_ALLOWED_CHAT_IDS", "").split(",")
                if item.strip()
            )
        except ValueError as exc:
            raise ValueError("TELEGRAM_ALLOWED_CHAT_IDS must contain integer IDs") from exc
        if require_telegram and (not token or not chats):
            raise ValueError("TELEGRAM_BOT_TOKEN and TELEGRAM_ALLOWED_CHAT_IDS are required")
        threshold = float(os.getenv("TELEGRAM_ALERT_THRESHOLD", "65"))
        top = int(os.getenv("TELEGRAM_TOP_LEADS", "5"))
        poll = int(os.getenv("TELEGRAM_POLL_TIMEOUT_SECONDS", "25"))
        if not 0 <= threshold <= 100:
            raise ValueError("TELEGRAM_ALERT_THRESHOLD must be 0..100")
        if not 1 <= top <= 20:
            raise ValueError("TELEGRAM_TOP_LEADS must be 1..20")
        if not 1 <= poll <= 50:
            raise ValueError("TELEGRAM_POLL_TIMEOUT_SECONDS must be 1..50")
        truthy = {"1", "true", "yes", "on"}
        return cls(
            token,
            chats,
            Path(os.getenv("FORENSICS_DB_PATH", ROOT / "forensics.db")).resolve(),
            Path(os.getenv("API_SNAPSHOT_PATH", runtime / "dashboard_candidates.json")).resolve(),
            Path(os.getenv("REPORT_OUTPUT_PATH", runtime / "phase1_fingerprint_report.json")).resolve(),
            Path(os.getenv("CANDIDATES_CSV_PATH", runtime / "phase1_fingerprint_report_candidates.csv")).resolve(),
            Path(os.getenv("TELEGRAM_DASHBOARD_PATH", runtime / "site-build" / "index.html")).resolve(),
            runtime,
            Path(os.getenv("TELEGRAM_HEALTH_PATH", runtime / "telegram_health.json")).resolve(),
            Path(os.getenv("TELEGRAM_STATE_PATH", runtime / "telegram_state.json")).resolve(),
            Path(os.getenv("TELEGRAM_LOCK_PATH", runtime / "telegram_bot.lock")).resolve(),
            threshold,
            top,
            poll,
            os.getenv("TELEGRAM_PUSH_ALERTS", "false").lower() in truthy,
            os.getenv("TELEGRAM_ALERT_EXISTING_ON_START", "false").lower() in truthy,
            int(os.getenv("TELEGRAM_HEALTH_MAX_AGE_SECONDS", "180")),
        )


class SingleInstance:
    def __init__(self, path):
        self.path = path
        self.handle = None

    def __enter__(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = open(self.path, "a+", encoding="utf-8")
        try:
            if os.name == "nt":
                import msvcrt
                self.handle.seek(0)
                self.handle.write(" ")
                self.handle.flush()
                self.handle.seek(0)
                msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            self.handle.close()
            raise RuntimeError("another Telegram bot instance owns the lock") from exc
        self.handle.seek(0)
        self.handle.truncate()
        self.handle.write(str(os.getpid()))
        self.handle.flush()
        return self

    def __exit__(self, *_):
        self.handle.close()


class Telegram:
    def __init__(self, token):
        self.url = "https://api.telegram.org/bot" + token
        self.http = requests.Session()

    def call(self, method, data=None, files=None, timeout=30):
        response = self.http.post(
            self.url + "/" + method, data=data, files=files, timeout=timeout
        )
        response.raise_for_status()
        payload = response.json()
        if not payload.get("ok"):
            raise RuntimeError(
                "Telegram " + method + ": " + payload.get("description", "unknown error")
            )
        return payload.get("result")

    def configure(self):
        self.call("deleteWebhook", {"drop_pending_updates": "false"})
        commands = [
            {"command": "leads", "description": "Top candidate leads"},
            {"command": "dashboard", "description": "Download interactive dashboard"},
            {"command": "refresh", "description": "Refresh DEX market observations"},
            {"command": "dbxlsx", "description": "Export formatted Excel workbook"},
            {"command": "dbcsv", "description": "Download raw candidate CSV"},
            {"command": "status", "description": "Collector and queue health"},
            {"command": "log", "description": "Audit recent collector and queue logs"},
            {"command": "ingest", "description": "Queue addresses from a forwarded token notice"},
            {"command": "sources", "description": "Manage intake & discovery sources"},
            {"command": "queue", "description": "View or safely restart unfinished token jobs"},
        ]
        self.call("setMyCommands", {"commands": json.dumps(commands)})

    def updates(self, offset, timeout):
        return self.call(
            "getUpdates",
            {
                "offset": offset,
                "timeout": timeout,
                "allowed_updates": json.dumps(["message", "callback_query"]),
            },
            timeout=timeout + 10,
        )

    def message(self, chat_id, body, reply_markup=None):
        data = {
            "chat_id": chat_id,
            "text": body,
            "parse_mode": "HTML",
            "link_preview_options": json.dumps({"is_disabled": True}),
        }
        if reply_markup is not None:
            data["reply_markup"] = json.dumps(reply_markup)
        self.call("sendMessage", data)

    def answer_callback(self, callback_id, text=""):
        self.call("answerCallbackQuery", {"callback_query_id": callback_id, "text": text})

    def document(self, chat_id, path, caption):
        with path.open("rb") as document:
            self.call(
                "sendDocument",
                {"chat_id": chat_id, "caption": caption},
                {"document": (path.name, document)},
                120,
            )


@contextmanager
def connect(path):
    connection = sqlite3.connect(path, timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA busy_timeout=30000")
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def ensure_schema(path):
    with connect(path) as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS telegram_alerts(
                chat_id INTEGER NOT NULL,
                chain_id INTEGER NOT NULL,
                ca TEXT NOT NULL,
                score REAL NOT NULL,
                sent_at TEXT NOT NULL,
                PRIMARY KEY(chat_id,chain_id,ca)
            );
            CREATE TABLE IF NOT EXISTS telegram_alert_baselines(
                chat_id INTEGER PRIMARY KEY,
                initialized_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_telegram_alerts_sent
                ON telegram_alerts(sent_at);
            CREATE TABLE IF NOT EXISTS forwarded_token_intake(
                intake_id INTEGER PRIMARY KEY AUTOINCREMENT,
                chat_id INTEGER NOT NULL,
                message_id INTEGER NOT NULL,
                ca TEXT NOT NULL,
                chain_id INTEGER,
                source_kind TEXT NOT NULL,
                source_text TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'held',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(chat_id, message_id, ca)
            );
            CREATE INDEX IF NOT EXISTS idx_forwarded_token_intake_status
                ON forwarded_token_intake(status, chain_id, created_at);
            """
        )


def enabled_chain_ids():
    try:
        return {int(value.strip()) for value in os.getenv("ENABLED_CHAIN_IDS", "4663,5042").split(",") if value.strip()}
    except ValueError as exc:
        raise ValueError("ENABLED_CHAIN_IDS must contain comma-separated integers") from exc


def queue_forwarded_notice(settings, chat_id, message):
    """Audit a forwarded notice and queue only configured supported chains and enabled sources."""
    text = forwarded_text(message)
    tokens = parse_forwarded_tokens(text)
    if not tokens:
        return {"queued": [], "held": [], "reason": "No EVM contract address found."}
    db = ForensicDatabase(str(settings.db))
    queue = QueueStore(db)
    enabled = enabled_chain_ids()
    queued, held, dev_seeds = [], [], []
    message_id = int(message.get("message_id") or 0)
    source_text = text[:8000]

    kind_to_switch = {
        "forwarded_pons_migration": "pons_forward",
        "forwarded_arc_token": "arc_forward",
        "forwarded_bsc_token": "bsc_forward",
    }

    for token in tokens:
        switch_key = kind_to_switch.get(token.source_kind)
        switch_enabled = is_source_enabled(switch_key, runtime_dir=settings.runtime) if switch_key else False
        chain_supported = token.chain_id in enabled and token.chain_id in CHAINS
        is_dev_seed = bool(token.metadata.get("is_dev_seed"))
        supported = (
            not is_dev_seed and chain_supported and switch_enabled
            and token.source_kind in {"forwarded_pons_migration", "forwarded_arc_token"}
        )
        status = "dev_seed" if is_dev_seed else ("queued" if supported else "held")
        with connect(settings.db) as connection:
            connection.execute(
                """INSERT INTO forwarded_token_intake(chat_id,message_id,ca,chain_id,source_kind,source_text,status)
                   VALUES(?,?,?,?,?,?,?)
                   ON CONFLICT(chat_id,message_id,ca) DO UPDATE SET
                     updated_at=CURRENT_TIMESTAMP, status=excluded.status""",
                (chat_id, message_id, token.ca, token.chain_id, token.source_kind, source_text, status),
            )
        if is_dev_seed:
            dev_seeds.append(token)
        elif supported:
            queue.enqueue(
                token.ca, token.source_kind,
                {
                    "forwarded_at": now(),
                    "forwarded_chat_id": chat_id,
                    "forwarded_message_id": message_id,
                    "source_kind": token.source_kind,
                    "source_text": source_text,
                    "dev_wallet": token.dev_wallet,
                    "pair_address": token.pair_address,
                    "symbol": token.symbol,
                    "name": token.name,
                    "metadata": token.metadata,
                },
                chain_id=token.chain_id, priority=5,
            )
            queued.append(token)
        else:
            held.append(token)
    return {"queued": queued, "held": held, "dev_seeds": dev_seeds, "reason": ""}

def load_candidates(settings):
    source = settings.snapshot if settings.snapshot.exists() else settings.report
    if not source.exists():
        raise FileNotFoundError("candidate snapshot missing: " + str(source))
    payload = json.loads(source.read_text(encoding="utf-8"))
    rows = payload.get("candidates") or payload.get("candidate_tokens") or []
    result = []
    for source_row in rows:
        row = dict(source_row)
        chain_id = int(row.get("chain_id") or 4663)
        row["chain_id"] = chain_id
        row["chain"] = CHAINS.get(chain_id, clean(row.get("chain"), str(chain_id)).upper())
        row["ca"] = clean(row.get("ca")).lower()
        row["symbol"] = clean(row.get("symbol"), "UNKNOWN")
        row["score"] = float(row.get("score", row.get("candidate_score", 0)) or 0)
        row["confidence"] = clean(row.get("confidence"), "UNRATED")
        row["is_qualified"] = bool(row.get("is_qualified", False))
        row["is_training_anchor"] = bool(row.get("is_training_anchor", False))
        row["is_dex_paid"] = bool(row.get("is_dex_paid", False))
        row["team"] = clean(
            row.get("team", row.get("inferred_team")), "Unclustered"
        )
        row["best_match_symbol"] = clean(row.get("best_match_symbol"), "Unknown")
        result.append(row)
    if settings.db.exists():
        query = """
            SELECT token_live_at,ath_usd,ath_source,current_market_cap_usd,
                   observed_peak_market_cap_usd,fdv_usd,current_liquidity_usd,market_data_at,
                   dev_wallet,funder_1hop,candidate_team,team_tier,
                   is_qualified,is_training_anchor,is_dex_paid
            FROM v_full_forensic_profile
            WHERE LOWER(ca)=? AND chain_id=? LIMIT 1
        """
        with connect(settings.db) as connection:
            for row in result:
                found = connection.execute(
                    query, (row["ca"], row["chain_id"])
                ).fetchone()
                if found:
                    for key in found.keys():
                        if found[key] is not None:
                            row[key] = found[key]
    return sorted(result, key=lambda item: item["score"], reverse=True)


def evidence_summary(row):
    parts = []
    for item in (row.get("evidence") or [])[:4]:
        parts.append(
            clean(item.get("feature"), "match")
            + ": "
            + clean(item.get("value"), "yes")
        )
    return "; ".join(parts) or "No evidence breakdown"


def lead_links(row):
    address = quote(row["ca"], safe="")
    chain_id = row["chain_id"]
    links = [
        '<a href="https://dexscreener.com/search?q='
        + address
        + '">DexScreener</a>'
    ]
    if chain_id == 4663:
        links.append(
            '<a href="https://gmgn.ai/robinhood/token/' + address + '">GMGN</a>'
        )
    if chain_id in EXPLORERS:
        links.append(
            '<a href="'
            + EXPLORERS[chain_id].format(ca=address)
            + '">Explorer</a>'
        )
    return " | ".join(links)


def lead_message(row, alert=False):
    symbol = html.escape(row["symbol"])
    chain = html.escape(row["chain"])
    team = html.escape(clean(row.get("candidate_team", row.get("team")), "Unclustered"))
    title = "<b>HIGH CONVICTION DEV ALERT</b>" if alert else "<b>$" + symbol + " [" + chain + "]</b>"
    token = "Token: <b>$" + symbol + "</b> [" + chain + "]\n" if alert else ""
    return (
        title + "\n" + token
        + "Confidence: <b>" + html.escape(row["confidence"])
        + "</b> - Score: <b>" + format(row["score"], ".1f")
        + "%</b>\nQualified: <b>" + ("YES" if row.get("is_qualified") else "NO")
        + "</b> (training anchor: " + ("YES" if row.get("is_training_anchor") else "NO")
        + ")\nCurrent MC: <b>$" + format(float(row.get("current_market_cap_usd", row.get("market_cap", 0)) or 0), ",.0f")
        + "</b> - Sourced ATH: <b>"
        + (("$" + format(float(row.get("ath_usd", row.get("ath", 0)) or 0), ",.0f")) if row.get("ath_source") else "N/A")
        + "</b>\nContract: <code>" + html.escape(row["ca"])
        + "</code>\nNearest sibling: $" + html.escape(row["best_match_symbol"])
        + " - Team: " + team
        + "\nDeployer: <code>" + html.escape(clean(row.get("dev_wallet"), "unknown"))
        + "</code>\n1-hop funder: <code>" + html.escape(clean(row.get("funder_1hop"), "unknown"))
        + "</code>\nLinks: " + lead_links(row)
    )


def export_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    headers = [
        "Symbol",
        "Chain",
        "Confidence",
        "Score",
        "Nearest Sibling",
        "Inferred Team",
        "Deployer",
        "1-Hop Funder",
        "Current MC (USD)",
        "Sourced ATH (USD)",
        "Launch Date (UTC)",
        "Top Evidence Matches",
        "Contract Address",
        "Rug",
        "Qualified",
        "Training Anchor",
    ]
    sheet.append(headers)
    blue = PatternFill("solid", fgColor="1F4E78")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    for cell in sheet[1]:
        cell.fill = blue
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        launched = row.get("token_live_at", row.get("token_live"))
        if launched:
            try:
                launched = datetime.fromisoformat(
                    str(launched).replace("Z", "+00:00")
                ).replace(tzinfo=None)
            except ValueError:
                launched = safe_cell(str(launched))
        rug = bool(row.get("is_rug")) or "RUG" in clean(
            row.get("team_tier")
        ).upper()
        sheet.append(
            [
                safe_cell(row["symbol"]),
                safe_cell(row["chain"]),
                safe_cell(row["confidence"]),
                row["score"] / 100,
                safe_cell(row["best_match_symbol"]),
                safe_cell(
                    clean(
                        row.get("candidate_team", row.get("team")),
                        "Unclustered",
                    )
                ),
                safe_cell(clean(row.get("dev_wallet"))),
                safe_cell(clean(row.get("funder_1hop"))),
                float(row.get("current_market_cap_usd", row.get("market_cap", 0)) or 0),
                float(row.get("ath_usd", row.get("ath", 0)) or 0) if row.get("ath_source") else None,
                launched,
                safe_cell(evidence_summary(row)),
                safe_cell(row["ca"]),
                "YES" if rug else "NO",
                "YES" if row.get("is_qualified") else "NO",
                "YES" if row.get("is_training_anchor") else "NO",
            ]
        )
        fill = red if rug else green if row["confidence"] == "HIGH_LEAD" else None
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
    sheet.freeze_panes = "A2"
    widths = [14, 10, 18, 12, 22, 22, 46, 46, 16, 16, 22, 60, 46, 10, 14, 18]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cell in sheet["D"][1:]:
        cell.number_format = "0.0%"
    for cell in sheet["I"][1:]:
        cell.number_format = "$#,##0.00"
    for column in ("G", "H", "L"):
        for cell in sheet[column][1:]:
            cell.number_format = "@"
            cell.quotePrefix = True
    for cell in sheet["J"][1:]:
        if isinstance(cell.value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
    sheet.auto_filter.ref = "A1:M" + str(max(1, sheet.max_row))
    if sheet.max_row > 1:
        table = Table(
            displayName="OmniRebornLeads",
            ref="A1:M" + str(sheet.max_row),
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    workbook.save(path)
    return path


def regenerate_dashboard(settings):
    environment = os.environ.copy()
    environment["DASHBOARD_OUTPUT_DIR"] = str(settings.dashboard.parent)
    environment["DASHBOARD_API_URL"] = ""
    subprocess.run(
        [sys.executable, str(ROOT / "generate_html_dashboard.py")],
        cwd=ROOT,
        env=environment,
        check=True,
        timeout=180,
    )


def source_switch_keyboard(switches):
    """Build authenticated Telegram callback controls for each source."""
    rows = []
    for key, definition in SOURCE_DEFINITIONS.items():
        enabled = bool(switches.get(key, False))
        next_state = "off" if enabled else "on"
        icon = "🟢 ON" if enabled else "⚪ OFF"
        rows.append([{
            "text": f"{definition['label']}: {icon}",
            "callback_data": f"source:{key}:{next_state}",
        }])
    return {"inline_keyboard": rows}


def source_switch_message(settings):
    switches = load_source_switches(runtime_dir=settings.runtime)
    lines = ["<b>⚙️ OmniReborn Intake & Discovery Sources</b>"]
    for key, definition in SOURCE_DEFINITIONS.items():
        icon = "🟢 ON" if switches.get(key, False) else "⚪ OFF"
        lines.append(f"• <b>{html.escape(definition['label'])}:</b> {icon}\n  <i>{html.escape(definition['description'])}</i>")
    lines.append("\n<i>Tap a button to switch it. Command fallback: /sources &lt;name&gt; [on|off]</i>")
    return "\n".join(lines), source_switch_keyboard(switches)

def status_message(settings, started, data_error=""):
    counts = {}
    latest = "none"
    if settings.db.exists():
        with connect(settings.db) as connection:
            counts = {row["status"]: row["n"] for row in connection.execute(
                "SELECT status,COUNT(*) n FROM ingestion_jobs GROUP BY status"
            )}
            row = connection.execute(
                "SELECT status,started_at FROM ingestion_runs ORDER BY run_id DESC LIMIT 1"
            ).fetchone()
            if row:
                latest = row["status"] + " (started " + row["started_at"] + ")"
    uptime = int(time.monotonic() - started)
    size_kb = settings.db.stat().st_size / 1024 if settings.db.exists() else 0
    data_state = "ready" if not data_error else "degraded: " + data_error[:180]
    return (
        "<b>OmniReborn Phase 1 status</b>\n"
        + "Bot uptime: " + str(uptime // 3600) + "h " + str((uptime % 3600) // 60) + "m\n"
        + "Database: " + format(size_kb, ".1f") + " KB\n"
        + "Queue: pending=" + str(counts.get("pending", 0))
        + ", retry=" + str(counts.get("retry", 0))
        + ", running=" + str(counts.get("running", 0))
        + ", dead=" + str(counts.get("dead", 0))
        + ", succeeded=" + str(counts.get("succeeded", 0))
        + "\nLatest collection: " + html.escape(latest)
        + "\nCandidate data: " + html.escape(data_state)
        + "\nChains: " + html.escape(os.getenv("ENABLED_CHAIN_IDS", "4663,5042"))
        + "\nDiscovery mode: " + html.escape(os.getenv("DISCOVERY_MODE", "forwarded_only"))
        + "\nActive sources: " + html.escape(
            ", ".join(
                SOURCE_DEFINITIONS[k]["label"]
                for k, v in load_source_switches(runtime_dir=settings.runtime).items()
                if v and k in SOURCE_DEFINITIONS
            ) or "None (idle)"
        )
        + "\nCollection cadence: " + html.escape(os.getenv("COLLECTION_INTERVAL_SECONDS", "300")) + " seconds"
        + "\nPhase 2 push alerts: " + ("enabled" if settings.push_alerts else "disabled")
    )

class BotService:
    def __init__(self, settings, api=None):
        self.settings = settings
        self.api = api or Telegram(settings.token)
        self.started = time.monotonic()
        self.offset = 0
        self.error = ""
        if settings.state.exists():
            try:
                self.offset = int(
                    json.loads(settings.state.read_text(encoding="utf-8")).get(
                        "offset", 0
                    )
                )
            except (OSError, ValueError, json.JSONDecodeError):
                LOG.warning("Ignoring invalid Telegram state")

    def write_health(self, state="healthy"):
        atomic_json(
            self.settings.health,
            {
                "status": state,
                "updated_at": now(),
                "unix_time": time.time(),
                "pid": os.getpid(),
                "offset": self.offset,
                "last_error": self.error,
            },
        )

    def initialize_baseline(self, rows):
        qualifying = [
            row for row in rows if row["score"] >= self.settings.threshold
        ]
        with connect(self.settings.db) as connection:
            for chat_id in self.settings.chats:
                exists = connection.execute(
                    "SELECT 1 FROM telegram_alert_baselines WHERE chat_id=?",
                    (chat_id,),
                ).fetchone()
                if exists:
                    continue
                if not self.settings.alert_existing:
                    connection.executemany(
                        "INSERT OR IGNORE INTO telegram_alerts VALUES(?,?,?,?,?)",
                        [
                            (
                                chat_id,
                                row["chain_id"],
                                row["ca"],
                                row["score"],
                                now(),
                            )
                            for row in qualifying
                        ],
                    )
                connection.execute(
                    "INSERT INTO telegram_alert_baselines VALUES(?,?)",
                    (chat_id, now()),
                )

    def send_new_alerts(self, rows):
        if not self.settings.push_alerts:
            qualifying = [row for row in rows if row["score"] >= self.settings.threshold]
            with connect(self.settings.db) as connection:
                for chat_id in self.settings.chats:
                    connection.executemany(
                        "INSERT OR IGNORE INTO telegram_alerts VALUES(?,?,?,?,?)",
                        [
                            (chat_id, row["chain_id"], row["ca"], row["score"], now())
                            for row in qualifying
                        ],
                    )
            return 0
        sent = 0
        for chat_id in self.settings.chats:
            for row in rows:
                if row["score"] < self.settings.threshold:
                    continue
                with connect(self.settings.db) as connection:
                    exists = connection.execute(
                        "SELECT 1 FROM telegram_alerts "
                        "WHERE chat_id=? AND chain_id=? AND ca=?",
                        (chat_id, row["chain_id"], row["ca"]),
                    ).fetchone()
                if exists:
                    continue
                self.api.message(chat_id, lead_message(row, True))
                with connect(self.settings.db) as connection:
                    connection.execute(
                        "INSERT OR IGNORE INTO telegram_alerts VALUES(?,?,?,?,?)",
                        (
                            chat_id,
                            row["chain_id"],
                            row["ca"],
                            row["score"],
                            now(),
                        ),
                    )
                sent += 1
        return sent

    def candidate_rows(self, context):
        try:
            rows = load_candidates(self.settings)
        except Exception as exc:
            self.error = "candidate data unavailable: " + f"{type(exc).__name__}: {exc}"[:420]
            LOG.warning("Candidate data unavailable during %s: %s", context, exc)
            return None
        self.error = ""
        return rows

    def command(self, chat_id, body, message=None):
        command = body.split()[0].split("@")[0].lower()
        if command == "/leads":
            rows = self.candidate_rows("/leads")
            if rows is None:
                raise RuntimeError(self.error)
            self.api.message(
                chat_id,
                "<b>Top Phase 1 candidate leads</b>\n\n"
                + "\n\n".join(lead_message(row) for row in rows[: self.settings.top]),
            )
        elif command == "/dashboard":
            try:
                regenerate_dashboard(self.settings)
            except Exception:
                LOG.exception("Dashboard refresh failed; using last good copy")
            if not self.settings.dashboard.exists():
                raise FileNotFoundError("dashboard unavailable")
            self.api.document(chat_id, self.settings.dashboard, "OmniReborn Phase 1 dashboard")
        elif command == "/dbxlsx":
            rows = self.candidate_rows("/dbxlsx")
            if rows is None:
                raise RuntimeError(self.error)
            path = export_xlsx(rows, self.settings.runtime / "OmniReborn_Leads.xlsx")
            self.api.document(chat_id, path, "OmniReborn forensic leads")
        elif command == "/dbcsv":
            path = self.settings.csv if self.settings.csv.exists() else ROOT / "phase1_fingerprint_report_candidates.csv"
            if not path.exists():
                raise FileNotFoundError("candidate CSV unavailable")
            self.api.document(chat_id, path, "OmniReborn raw candidate data")
        elif command in ("/ingest", "/intake"):
            result = queue_forwarded_notice(self.settings, chat_id, message or {"text": body})
            if result["reason"]:
                self.api.message(chat_id, "⚠️ " + html.escape(result["reason"]))
            else:
                queued = result["queued"]
                held = result["held"]
                dev_seeds = result.get("dev_seeds", [])
                lines = ["<b>Forwarded source intake</b>"]
                if queued:
                    for item in queued:
                        tag = CHAINS.get(item.chain_id, str(item.chain_id))
                        extra = []
                        if item.symbol:
                            extra.append(f"${item.symbol}")
                        if item.dev_wallet:
                            extra.append(f"Dev: <code>{item.dev_wallet[:8]}...</code>")
                        extra_str = f" ({', '.join(extra)})" if extra else ""
                        lines.append(f"Queued for full forensic verification: <code>{html.escape(item.ca)}</code> [{tag}]{extra_str}")
                if dev_seeds:
                    for item in dev_seeds:
                        lines.append(
                            "Recorded developer seed (no token job queued): <code>"
                            + html.escape(item.dev_wallet or item.ca) + "</code>"
                        )
                if held:
                    for item in held:
                        tag = str(item.chain_id or "unknown")
                        extra = []
                        if item.dev_wallet:
                            extra.append(f"Dev: <code>{item.dev_wallet[:8]}...</code>")
                        extra_str = f" ({', '.join(extra)})" if extra else ""
                        lines.append(f"Held without enrichment (disabled/unsupported): <code>{html.escape(item.ca)}</code> [{tag}]{extra_str}")
                lines.append("Forwarded text is intake evidence only; qualification still requires chain verification.")
                self.api.message(chat_id, "\n".join(lines))
        elif command in ("/queue", "/queue_restart"):
            action = body.split(maxsplit=1)[1].strip().lower() if len(body.split(maxsplit=1)) > 1 else "status"
            queue = QueueStore(ForensicDatabase(str(self.settings.db)))
            if action in {"restart", "reset"}:
                restarted = queue.restart_incomplete()
                self.api.message(
                    chat_id,
                    "✅ Restarted <b>" + str(restarted) + "</b> unfinished token job(s). "
                    "Completed forensic records were preserved.",
                )
            elif action in {"status", "list"}:
                stats = queue.stats()
                self.api.message(
                    chat_id,
                    "<b>Token queue</b>\n"
                    + "pending=" + str(stats["pending"])
                    + ", retry=" + str(stats["retry"])
                    + ", running=" + str(stats["running"])
                    + ", dead=" + str(stats["dead"])
                    + ", succeeded=" + str(stats["succeeded"])
                    + "\nUse <code>/queue restart</code> to restart only unfinished jobs.",
                )
            else:
                self.api.message(chat_id, "Usage: <code>/queue</code> or <code>/queue restart</code>")
        elif command in ("/sources", "/source"):
            parts = body.split()
            if len(parts) >= 2:
                target_source = normalize_source_name(parts[1])
                if not target_source:
                    avail = ", ".join(SOURCE_DEFINITIONS.keys())
                    self.api.message(chat_id, f"⚠️ Unknown source <code>{html.escape(parts[1])}</code>.\nAvailable: {avail}")
                    return
                if len(parts) >= 3:
                    new_state = parts[2].strip().lower() in {"1", "true", "yes", "on", "enable"}
                else:
                    curr = is_source_enabled(target_source, runtime_dir=self.settings.runtime)
                    new_state = not curr
                set_source_switch(target_source, new_state, runtime_dir=self.settings.runtime)
                icon = "🟢 ON" if new_state else "⚪ OFF"
                label = SOURCE_DEFINITIONS[target_source]["label"]
                self.api.message(chat_id, f"✅ <b>{html.escape(label)}</b> is now <b>{icon}</b>")
                return

            source_body, keyboard = source_switch_message(self.settings)
            self.api.message(chat_id, source_body, reply_markup=keyboard)
        elif command == "/status":
            self.api.message(chat_id, status_message(self.settings, self.started, self.error))
        elif command in ("/log", "/logs"):
            log_path = self.settings.runtime / "streamer.log"
            content = ""
            if log_path.exists():
                try:
                    lines = log_path.read_text(encoding="utf-8", errors="replace").splitlines()
                    content = "\n".join(lines[-45:])
                except Exception as exc:
                    content = f"Error reading log: {exc}"
            if not content:
                content = "No log entries found in streamer.log yet."
            if len(content) <= 3800:
                self.api.message(chat_id, f"📋 <b>Recent Collector Logs:</b>\n<pre>{html.escape(content)}</pre>")
            else:
                audit_file = self.settings.runtime / "audit_recent.log"
                audit_file.write_text(content, encoding="utf-8")
                self.api.document(chat_id, audit_file, "OmniReborn audit logs")
        elif command in ("/refresh", "/refreshmarket"):
            self.api.message(chat_id, "⏳ Refreshing DEX market data for all tokens...")
            try:
                refresh_env = os.environ.copy()
                subprocess.run(
                    [sys.executable, str(ROOT / "refresh_market_data.py"), "--db", str(self.settings.db), "--all"],
                    cwd=ROOT,
                    env=refresh_env,
                    check=True,
                    timeout=120,
                )
                regenerate_dashboard(self.settings)
                self.api.document(chat_id, self.settings.dashboard, "OmniReborn Phase 1 dashboard (Refreshed)")
            except Exception as exc:
                LOG.exception("Market refresh failed")
                self.api.message(chat_id, "⚠️ Market refresh failed: " + html.escape(str(exc)))
        else:
            self.api.message(chat_id, "Commands: /leads /dashboard /refresh /dbxlsx /dbcsv /ingest /sources /queue /status /log")
    def handle_source_callback(self, callback):
        message = callback.get("message") or {}
        chat_id = (message.get("chat") or {}).get("id")
        if chat_id not in self.settings.chats:
            LOG.warning("Ignored unauthorized source callback chat_id=%s", chat_id)
            return
        parts = str(callback.get("data") or "").split(":")
        if len(parts) != 3 or parts[0] != "source" or parts[2] not in {"on", "off"}:
            self.api.answer_callback(callback.get("id", ""), "Invalid source control")
            return
        source_key = normalize_source_name(parts[1])
        if not source_key:
            self.api.answer_callback(callback.get("id", ""), "Unknown source")
            return
        enabled = parts[2] == "on"
        set_source_switch(source_key, enabled, runtime_dir=self.settings.runtime)
        self.api.answer_callback(callback.get("id", ""), "Source updated")
        body, keyboard = source_switch_message(self.settings)
        self.api.message(int(chat_id), body, reply_markup=keyboard)

    def handle_update(self, update):
        callback = update.get("callback_query") or {}
        if callback:
            try:
                self.handle_source_callback(callback)
            except Exception as exc:
                LOG.exception("Source callback failed")
                self.api.answer_callback(callback.get("id", ""), "Update failed safely")
            return
        message = update.get("message") or {}
        chat_id = (message.get("chat") or {}).get("id")
        body = clean(message.get("text"))
        if chat_id not in self.settings.chats:
            LOG.warning("Ignored unauthorized chat_id=%s", chat_id)
            return
        if body.startswith("/"):
            try:
                self.command(int(chat_id), body, message)
            except Exception as exc:
                LOG.exception("Command failed")
                self.api.message(
                    int(chat_id),
                    "⚠️ Command failed safely: " + html.escape(str(exc)),
                )
        elif is_forwarded_message(message) and parse_forwarded_tokens(forwarded_text(message)):
            # The production bot is the sole long-poll listener. Forwarding a
            # recognized notice into the authorized group ingests it directly.
            try:
                self.command(int(chat_id), "/ingest", message)
            except Exception as exc:
                LOG.exception("Automatic forwarded intake failed")
                self.api.message(
                    int(chat_id),
                    "⚠️ Forwarded intake failed safely: " + html.escape(str(exc)),
                )

    def cycle(self):
        rows = self.candidate_rows("background alert cycle")
        if rows is not None:
            self.send_new_alerts(rows)
        # Poll commands even if reporting artifacts are unavailable or a dashboard build is in progress.
        for update in self.api.updates(
            self.offset, self.settings.poll
        ):
            self.handle_update(update)
            self.offset = max(
                self.offset, int(update["update_id"]) + 1
            )
            atomic_json(
                self.settings.state,
                {"offset": self.offset, "updated_at": now()},
            )
        self.write_health("healthy" if not self.error else "degraded")

    def run(self, once=False):
        with SingleInstance(self.settings.lock):
            ensure_schema(self.settings.db)
            self.api.configure()
            rows = self.candidate_rows("startup baseline")
            if rows is not None:
                self.initialize_baseline(rows)
            self.write_health("starting" if rows is not None else "degraded")
            backoff = 1
            while True:
                try:
                    self.cycle()
                    backoff = 1
                    if once:
                        return
                except KeyboardInterrupt:
                    return
                except Exception as exc:
                    self.error = str(exc)[:500]
                    LOG.exception(
                        "Bot cycle failed; retrying in %ss", backoff
                    )
                    self.write_health("degraded")
                    if once:
                        raise
                    time.sleep(backoff)
                    backoff = min(backoff * 2, 60)

def readiness(settings, online=True):
    missing = [
        str(path)
        for path in (settings.db, settings.snapshot)
        if not path.exists()
    ]
    if missing:
        raise RuntimeError("missing runtime files: " + ", ".join(missing))
    ensure_schema(settings.db)
    load_candidates(settings)
    if online:
        api = Telegram(settings.token)
        api.call("getMe")
        for chat_id in settings.chats:
            api.call("getChat", {"chat_id": chat_id})


def check_health(settings):
    if not settings.health.exists():
        raise RuntimeError("Telegram health file missing")
    payload = json.loads(settings.health.read_text(encoding="utf-8"))
    age = time.time() - float(payload.get("unix_time", 0))
    if age > settings.max_age:
        raise RuntimeError(
            "Telegram heartbeat stale (" + format(age, ".0f") + "s)"
        )


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--offline-check", action="store_true")
    parser.add_argument("--health-check", action="store_true")
    parser.add_argument("--once", action="store_true")
    parser.add_argument("--export-xlsx", type=Path)
    args = parser.parse_args(argv)
    logging.basicConfig(
        level=os.getenv("LOG_LEVEL", "INFO"),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )
    try:
        settings = Settings.from_env(
            not (args.export_xlsx or args.offline_check)
        )
        if args.export_xlsx:
            export_xlsx(
                load_candidates(settings), args.export_xlsx.resolve()
            )
            print(args.export_xlsx.resolve())
        elif args.health_check:
            check_health(settings)
        elif args.check or args.offline_check:
            readiness(settings, args.check)
            print("Telegram bot readiness check passed")
        else:
            BotService(settings).run(args.once)
        return 0
    except Exception as exc:
        LOG.error("%s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())