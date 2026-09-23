import json
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from openpyxl import load_workbook

from database import ForensicDatabase
from telegram_bot import (
    BotService,
    Settings,
    ensure_schema,
    export_xlsx,
    lead_links,
    load_candidates,
)


class FakeTelegram:
    def __init__(self):
        self.messages = []
        self.documents = []

    def message(self, chat_id, body):
        self.messages.append((chat_id, body))

    def document(self, chat_id, path, caption):
        self.documents.append((chat_id, path, caption))

    def configure(self):
        pass

    def updates(self, offset, timeout):
        return []


class TelegramBotTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.db_path = root / "forensics.db"
        ForensicDatabase(str(self.db_path))
        self.snapshot = root / "dashboard_candidates.json"
        self.snapshot.write_text(
            json.dumps(
                {
                    "candidates": [
                        {
                            "ca": "0x" + "a" * 40,
                            "chain_id": 5042,
                            "chain": "untrusted",
                            "symbol": "=FORMULA",
                            "score": 90.0,
                            "confidence": "HIGH_LEAD",
                            "team": "team arc",
                            "best_match_symbol": "ARCONE",
                            "evidence": [{"feature": "funder_1hop", "value": "0xabc"}],
                            "is_rug": False,
                        },
                        {
                            "ca": "0x" + "b" * 40,
                            "chain_id": 4663,
                            "symbol": "RUGGY",
                            "score": 70.0,
                            "confidence": "PROBABLE_LEAD",
                            "team": "team rug",
                            "best_match_symbol": "OLD",
                            "evidence": [],
                            "is_rug": True,
                        },
                    ]
                }
            ),
            encoding="utf-8",
        )
        self.settings = Settings(
            token="test",
            chats=(-1001,),
            db=self.db_path,
            snapshot=self.snapshot,
            report=root / "report.json",
            csv=root / "candidates.csv",
            dashboard=root / "index.html",
            runtime=root,
            health=root / "health.json",
            state=root / "state.json",
            lock=root / "bot.lock",
            push_alerts=False,
        )
        ensure_schema(self.db_path)

    def tearDown(self):
        self.temp.cleanup()

    def test_normalizes_chain_and_chain_specific_links(self):
        rows = load_candidates(self.settings)
        self.assertEqual(rows[0]["chain"], "ARC")
        self.assertNotIn("GMGN", lead_links(rows[0]))
        self.assertIn("https://gmgn.ai/robinhood/token/", lead_links(rows[1]))

    def test_unauthorized_chat_is_ignored(self):
        fake = FakeTelegram()
        service = BotService(self.settings, fake)
        service.handle_update(
            {"update_id": 1, "message": {"chat": {"id": 999}, "text": "/leads"}}
        )
        self.assertEqual(fake.messages, [])

    def test_phase1_disables_push_alerts_and_baseline_is_durable(self):
        rows = load_candidates(self.settings)
        fake = FakeTelegram()
        service = BotService(self.settings, fake)
        service.initialize_baseline(rows)
        later = dict(rows[0])
        later["ca"] = "0x" + "c" * 40
        self.assertEqual(service.send_new_alerts(rows + [later]), 0)
        enabled = Settings(**{**self.settings.__dict__, "push_alerts": True})
        self.assertEqual(BotService(enabled, fake).send_new_alerts(rows + [later]), 0)
        connection = sqlite3.connect(self.db_path)
        try:
            baseline = connection.execute(
                "SELECT COUNT(*) FROM telegram_alert_baselines"
            ).fetchone()[0]
            sent = connection.execute("SELECT COUNT(*) FROM telegram_alerts").fetchone()[0]
        finally:
            connection.close()
        self.assertEqual(baseline, 1)
        self.assertEqual(sent, 3)

    def test_phase2_push_is_deduplicated(self):
        enabled = Settings(**{**self.settings.__dict__, "push_alerts": True, "alert_existing": True})
        fake = FakeTelegram()
        service = BotService(enabled, fake)
        rows = load_candidates(enabled)
        service.initialize_baseline(rows)
        self.assertEqual(service.send_new_alerts(rows), 2)
        self.assertEqual(service.send_new_alerts(rows), 0)
        self.assertEqual(len(fake.messages), 2)

    def test_xlsx_is_typed_styled_and_formula_safe(self):
        output = self.settings.runtime / "OmniReborn_Leads.xlsx"
        export_xlsx(load_candidates(self.settings), output)
        workbook = load_workbook(output)
        sheet = workbook["Leads"]
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["D2"].data_type, "n")
        self.assertEqual(sheet["D2"].number_format, "0.0%")
        self.assertTrue(str(sheet["A2"].value).startswith("'="))
        self.assertEqual(sheet["A2"].fill.fgColor.rgb[-6:], "C6EFCE")
        self.assertEqual(sheet["A3"].fill.fgColor.rgb[-6:], "FFC7CE")
        self.assertIn("OmniRebornLeads", sheet.tables)

    def test_status_remains_responsive_when_candidate_data_is_unavailable(self):
        fake = FakeTelegram()
        fake.updates = lambda offset, timeout: [
            {"update_id": 1, "message": {"chat": {"id": -1001}, "text": "/status"}}
        ]
        service = BotService(self.settings, fake)
        with patch("telegram_bot.load_candidates", side_effect=FileNotFoundError("snapshot missing")):
            service.cycle()
        self.assertEqual(len(fake.messages), 1)
        self.assertIn("Candidate data: degraded", fake.messages[0][1])
        health = json.loads(self.settings.health.read_text(encoding="utf-8"))
        self.assertEqual(health["status"], "degraded")
    def test_refresh_command_dispatches_cleanly(self):
        fake = FakeTelegram()
        service = BotService(self.settings, fake)
        # Verify unknown command shows help with /refresh
        service.command(12345, "/unknown")
        self.assertIn("/refresh", fake.messages[-1][1])


if __name__ == "__main__":
    unittest.main()
